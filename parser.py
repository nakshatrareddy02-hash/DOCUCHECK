"""Document parsing: PDF, DOCX, TXT and (optionally OCR'd) images."""
from __future__ import annotations
import io
import os

ALLOWED = {"pdf", "docx", "txt", "jpg", "jpeg", "png"}
MAX_MB = 25


def human_size(n: float) -> str:
    if n < 1024:
        return f"{int(n)} B"
    for unit in ("KB", "MB", "GB"):
        n /= 1024
        if n < 1024 or unit == "GB":
            return f"{n:.1f} {unit}"


def _empty(name, data, ext):
    return {"id": f"{name}:{len(data)}", "name": name, "ext": ext.upper(), "size": len(data),
            "size_label": human_size(len(data)), "pages": 1, "text": "", "images": 0,
            "headings": [], "ocr": False, "error": None, "warning": None}


def _pdf(data, doc):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            pass
    doc["pages"] = len(reader.pages)
    parts, imgs = [], 0
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
        try:
            imgs += len(page.images)
        except Exception:
            pass
    doc["text"], doc["images"] = "\n".join(parts), imgs


def _docx(data, doc):
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
            if p.style is not None and p.style.name.lower().startswith(("heading", "title")):
                doc["headings"].append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    doc["text"] = "\n".join(lines)
    doc["images"] = sum(1 for r in d.part.rels.values() if "image" in r.reltype)
    doc["pages"] = max(1, round(len(doc["text"].split()) / 500))


def _txt(data, doc):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            doc["text"] = data.decode(enc)
            break
        except Exception:
            continue
    doc["pages"] = max(1, round(len(doc["text"].split()) / 500))


def _image(data, doc):
    doc["images"] = 1
    try:
        from PIL import Image
        import pytesseract
        doc["text"] = pytesseract.image_to_string(Image.open(io.BytesIO(data)))
        doc["ocr"] = True
    except ImportError:
        doc["warning"] = "OCR packages are not installed, so no text could be read from this image."
    except Exception as exc:  # tesseract binary missing, bad image ...
        doc["warning"] = ("OCR is unavailable (install the Tesseract engine to read scanned images): "
                          f"{type(exc).__name__}")


def parse_document(name: str, data: bytes) -> dict:
    """Parse an uploaded file. Never raises; problems are reported in doc['error']."""
    ext = os.path.splitext(name)[1].lower().lstrip(".")
    doc = _empty(name, data, ext)
    if ext not in ALLOWED:
        doc["error"] = f"Unsupported file type '.{ext}'. Please upload PDF, DOCX, TXT, JPG or PNG."
        return doc
    if len(data) == 0:
        doc["error"] = "The file is empty."
        return doc
    if len(data) > MAX_MB * 1024 * 1024:
        doc["error"] = f"File is larger than {MAX_MB} MB."
        return doc
    try:
        {"pdf": _pdf, "docx": _docx, "txt": _txt}.get(ext, _image)(data, doc)
    except ImportError as exc:
        doc["error"] = f"A required package is missing: {exc}. Run 'pip install -r requirements.txt'."
    except Exception as exc:
        doc["error"] = f"The file could not be read ({type(exc).__name__}). It may be corrupted or invalid."
    if not doc["error"] and not doc["text"].strip():
        if ext in ("jpg", "jpeg", "png"):
            doc["error"] = doc["warning"] or "No readable text was found in this image."
        else:
            doc["error"] = ("No text could be extracted. The document may be empty or a scanned PDF "
                            "(upload scanned pages as JPG/PNG so OCR can read them).")
    return doc
